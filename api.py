import io
import json
import os
import re
from typing import Any

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader


# ============================================================
# Configuration
# ============================================================

APP_TITLE = "ATS Resume Analyzer"
MODEL_NAME = "gemini-2.5-flash"

MAX_FILE_SIZE_MB = 10
ALLOWED_EXTENSIONS = ["pdf", "docx"]


# ============================================================
# Page configuration
# ============================================================

st.set_page_config(
    page_title=APP_TITLE,
    page_icon="📄",
    layout="wide",
)


# ============================================================
# Custom styling
# ============================================================

st.markdown(
    """
    <style>
        .main-title {
            font-size: 2.5rem;
            font-weight: 700;
            margin-bottom: 0.25rem;
        }

        .subtitle {
            color: #6b7280;
            font-size: 1.05rem;
            margin-bottom: 2rem;
        }

        .score-box {
            padding: 1.5rem;
            border-radius: 15px;
            text-align: center;
            background: linear-gradient(
                135deg,
                #f8fafc 0%,
                #eef2ff 100%
            );
            border: 1px solid #e5e7eb;
        }

        .score-number {
            font-size: 4rem;
            font-weight: 800;
            line-height: 1;
        }

        .score-label {
            color: #6b7280;
            font-size: 1rem;
        }

        .good {
            color: #16a34a;
        }

        .medium {
            color: #ca8a04;
        }

        .poor {
            color: #dc2626;
        }

        .info-card {
            padding: 1rem;
            border-radius: 10px;
            border: 1px solid #e5e7eb;
            background: #ffffff;
            margin-bottom: 0.75rem;
        }
    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# Utility functions
# ============================================================

def get_api_key() -> str:
    """
    Get Gemini API key from Streamlit secrets first,
    then environment variables.
    """
    try:
        secret_key = st.secrets.get("GEMINI_API_KEY")
        if secret_key:
            return secret_key
    except Exception:
        pass

    env_key = os.getenv("GEMINI_API_KEY")
    if env_key:
        return env_key

    return ""


def extract_pdf_text(file_bytes: bytes) -> str:
    """Extract text from a PDF."""
    reader = PdfReader(io.BytesIO(file_bytes))

    pages = []

    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)

    return "\n\n".join(pages).strip()


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract text from a DOCX file."""
    document = Document(io.BytesIO(file_bytes))

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    # Also collect table text because resumes sometimes use tables.
    table_text = []

    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            )

            if row_text:
                table_text.append(row_text)

    all_text = paragraphs + table_text

    return "\n".join(all_text).strip()


def extract_resume_text(uploaded_file) -> str:
    """Extract readable text from PDF or DOCX."""
    file_bytes = uploaded_file.getvalue()

    extension = uploaded_file.name.lower().split(".")[-1]

    if extension == "pdf":
        return extract_pdf_text(file_bytes)

    if extension == "docx":
        return extract_docx_text(file_bytes)

    raise ValueError(
        "Unsupported file format. Please upload a PDF or DOCX resume."
    )


def normalize_score(score: Any) -> int:
    """Safely convert a model score into 0-100."""
    try:
        score = int(float(score))
    except (TypeError, ValueError):
        score = 0

    return max(0, min(100, score))


def score_color(score: int) -> str:
    if score >= 80:
        return "good"

    if score >= 60:
        return "medium"

    return "poor"


def clean_json_response(text: str) -> dict:
    """
    Safely parse JSON returned by Gemini.

    Handles cases where the model accidentally wraps JSON
    in markdown code fences.
    """
    text = text.strip()

    # Remove markdown fences if present.
    text = re.sub(r"^```json\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"^```\s*", "", text)
    text = re.sub(r"\s*```$", "", text)

    try:
        return json.loads(text)
    except json.JSONDecodeError:
        # Attempt to find the outermost JSON object.
        start = text.find("{")
        end = text.rfind("}")

        if start != -1 and end != -1 and end > start:
            return json.loads(text[start:end + 1])

        raise ValueError("Gemini returned invalid JSON.")


# ============================================================
# Gemini analysis
# ============================================================

def analyze_resume(resume_text: str, job_description: str = "") -> dict:
    """
    Analyze a resume using Gemini 2.5 Flash.

    The ATS score is an ATS-readiness estimate, not a score
    generated by a particular ATS vendor.
    """
    api_key = get_api_key()

    if not api_key:
        raise RuntimeError(
            "Gemini API key not found. Add GEMINI_API_KEY to "
            "Streamlit secrets or your environment variables."
        )

    client = genai.Client(api_key=api_key)

    job_context = (
        job_description.strip()
        if job_description.strip()
        else "No job description was provided. Evaluate the resume for general ATS readiness."
    )

    prompt = f"""
You are an expert resume reviewer, recruiter, and ATS optimization specialist.

Analyze the resume below.

IMPORTANT:
- Do NOT claim this is the exact score from Workday, Greenhouse, Lever,
  Taleo, or another proprietary ATS.
- Produce an "ATS readiness score" from 0 to 100 based on the criteria below.
- Be practical and evidence-based.
- Do not invent experience, qualifications, achievements, or skills.
- Suggestions must preserve the candidate's truthfulness.
- Focus on improvements that genuinely help resume parsing and recruiter review.

ATS READINESS CRITERIA:

1. Contact information: 10 points
2. Standard resume sections: 10 points
3. Keyword and skill relevance: 20 points
4. Work experience quality: 15 points
5. Quantified achievements/results: 10 points
6. Formatting and ATS parseability: 15 points
7. Education/certifications: 5 points
8. Clarity, conciseness, and professional language: 10 points
9. Overall consistency: 5 points

TOTAL: 100 points.

JOB DESCRIPTION:
{job_context}

RESUME:
----------------
{resume_text}
----------------

Return ONLY valid JSON matching this exact structure:

{{
  "ats_score": 0,
  "score_summary": "Short explanation of the score",
  "category_scores": [
    {{
      "category": "Contact information",
      "score": 0,
      "max_score": 10,
      "feedback": "Specific feedback"
    }},
    {{
      "category": "Standard resume sections",
      "score": 0,
      "max_score": 10,
      "feedback": "Specific feedback"
    }},
    {{
      "category": "Keyword and skill relevance",
      "score": 0,
      "max_score": 20,
      "feedback": "Specific feedback"
    }},
    {{
      "category": "Work experience quality",
      "score": 0,
      "max_score": 15,
      "feedback": "Specific feedback"
    }},
    {{
      "category": "Quantified achievements/results",
      "score": 0,
      "max_score": 10,
      "feedback": "Specific feedback"
    }},
    {{
      "category": "Formatting and ATS parseability",
      "score": 0,
      "max_score": 15,
      "feedback": "Specific feedback"
    }},
    {{
      "category": "Education/certifications",
      "score": 0,
      "max_score": 5,
      "feedback": "Specific feedback"
    }},
    {{
      "category": "Clarity and professional language",
      "score": 0,
      "max_score": 10,
      "feedback": "Specific feedback"
    }},
    {{
      "category": "Overall consistency",
      "score": 0,
      "max_score": 5,
      "feedback": "Specific feedback"
    }}
  ],
  "strengths": [
    "Strength 1",
    "Strength 2",
    "Strength 3"
  ],
  "critical_issues": [
    "Issue that should be fixed first"
  ],
  "improvements": [
    {{
      "priority": "High",
      "area": "Area name",
      "problem": "What is wrong",
      "recommendation": "What the candidate should do",
      "example": "Example of an improved version if possible"
    }}
  ],
  "missing_sections": [
    "Section name"
  ],
  "keyword_analysis": {{
    "important_keywords_found": [
      "keyword"
    ],
    "important_keywords_missing": [
      "keyword"
    ],
    "note": "Explain the keyword analysis"
  }},
  "bullet_point_improvements": [
    {{
      "original": "Existing bullet from resume",
      "improved": "Truth-preserving improved bullet",
      "reason": "Why the new version is better"
    }}
  ],
  "formatting_checklist": [
    {{
      "item": "Formatting item",
      "status": "Good",
      "recommendation": "Recommendation"
    }}
  ],
  "final_recommendation": "Short final recommendation"
}}

Rules:
- ats_score must be an integer from 0 to 100.
- Category scores must not exceed their max_score.
- Do not invent facts.
- If something is unknown, say so.
- Keep feedback specific.
"""

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )

    if not response.text:
        raise RuntimeError("Gemini returned an empty response.")

    result = clean_json_response(response.text)

    # Normalize important fields.
    result["ats_score"] = normalize_score(
        result.get("ats_score", 0)
    )

    result.setdefault("score_summary", "")
    result.setdefault("category_scores", [])
    result.setdefault("strengths", [])
    result.setdefault("critical_issues", [])
    result.setdefault("improvements", [])
    result.setdefault("missing_sections", [])
    result.setdefault(
        "keyword_analysis",
        {
            "important_keywords_found": [],
            "important_keywords_missing": [],
            "note": "",
        },
    )
    result.setdefault("bullet_point_improvements", [])
    result.setdefault("formatting_checklist", [])
    result.setdefault("final_recommendation", "")

    return result


# ============================================================
# UI rendering
# ============================================================

def render_score(score: int):
    color = score_color(score)

    if score >= 80:
        label = "Strong ATS readiness"
    elif score >= 60:
        label = "Needs improvement"
    else:
        label = "Significant improvement needed"

    st.markdown(
        f"""
        <div class="score-box">
            <div class="score-number {color}">{score}/100</div>
            <div class="score-label">{label}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_category_scores(category_scores: list):
    if not category_scores:
        return

    st.subheader("📊 Score Breakdown")

    for item in category_scores:
        category = item.get("category", "Unknown")
        score = normalize_score(item.get("score", 0))
        max_score = normalize_score(item.get("max_score", 100))
        feedback = item.get("feedback", "")

        if max_score > 0:
            percentage = min(100, int((score / max_score) * 100))
        else:
            percentage = 0

        st.markdown(f"**{category} — {score}/{max_score}**")
        st.progress(percentage)
        st.caption(feedback)


def render_improvements(improvements: list):
    if not improvements:
        st.info("No specific improvements were returned.")
        return

    for index, item in enumerate(improvements, start=1):
        priority = item.get("priority", "Medium")
        area = item.get("area", "General")
        problem = item.get("problem", "")
        recommendation = item.get("recommendation", "")
        example = item.get("example", "")

        with st.expander(
            f"{index}. {priority} priority — {area}"
        ):
            st.markdown(f"**Problem:** {problem}")
            st.markdown(f"**Recommendation:** {recommendation}")

            if example:
                st.markdown("**Example:**")
                st.info(example)


# ============================================================
# Main application
# ============================================================

st.markdown(
    '<div class="main-title">📄 ATS Resume Analyzer</div>',
    unsafe_allow_html=True,
)

st.markdown(
    """
    <div class="subtitle">
    Upload your resume and get an AI-powered ATS-readiness score,
    keyword analysis, and actionable improvements.
    </div>
    """,
    unsafe_allow_html=True,
)

with st.sidebar:
    st.header("⚙️ Settings")

    st.markdown(
        """
        **Supported files**
        - PDF
        - DOCX

        **AI model**
        - Gemini 2.5 Flash

        **Score**
        - 0–100 ATS readiness estimate
        """
    )

    st.divider()

    st.caption(
        "Your resume is analyzed by Gemini. Do not upload confidential "
        "information unless you are comfortable sending it to the "
        "configured AI service."
    )


uploaded_file = st.file_uploader(
    "Upload your resume",
    type=ALLOWED_EXTENSIONS,
    help="Maximum recommended file size: 10 MB.",
)

job_description = st.text_area(
    "Optional: Paste the job description",
    height=220,
    placeholder=(
        "Paste the job description here to get "
        "job-specific keyword and ATS recommendations..."
    ),
)

analyze_button = st.button(
    "🔍 Analyze Resume",
    type="primary",
    use_container_width=True,
)

if analyze_button:

    if uploaded_file is None:
        st.warning("Please upload a PDF or DOCX resume first.")
        st.stop()

    file_size_mb = uploaded_file.size / (1024 * 1024)

    if file_size_mb > MAX_FILE_SIZE_MB:
        st.error(
            f"File is too large ({file_size_mb:.1f} MB). "
            f"Please upload a file smaller than {MAX_FILE_SIZE_MB} MB."
        )
        st.stop()

    try:
        with st.spinner("Reading your resume..."):
            resume_text = extract_resume_text(uploaded_file)

        if not resume_text.strip():
            st.error(
                "I could not extract readable text from this resume. "
                "If it is a scanned/image-only PDF, please use a text-based "
                "PDF or DOCX version."
            )
            st.stop()

        # Prevent sending an accidentally enormous extracted document.
        # This is only a safety guard; normal resumes will be much smaller.
        if len(resume_text) > 120_000:
            resume_text = resume_text[:120_000]

        with st.spinner("Gemini is analyzing your resume..."):
            result = analyze_resume(
                resume_text=resume_text,
                job_description=job_description,
            )

        st.session_state["analysis_result"] = result
        st.session_state["resume_name"] = uploaded_file.name

    except Exception as exc:
        st.error(
            "Something went wrong while analyzing the resume."
        )

        with st.expander("Technical details"):
            st.exception(exc)


# ============================================================
# Results
# ============================================================

if "analysis_result" in st.session_state:

    result = st.session_state["analysis_result"]

    st.divider()

    st.subheader(
        f"Results for `{st.session_state.get('resume_name', 'Resume')}`"
    )

    col1, col2 = st.columns([1, 2])

    with col1:
        render_score(result["ats_score"])

    with col2:
        st.markdown("### 📝 Summary")
        st.write(result.get("score_summary", ""))

        final_recommendation = result.get(
            "final_recommendation",
            "",
        )

        if final_recommendation:
            st.info(final_recommendation)

    st.divider()

    render_category_scores(
        result.get("category_scores", [])
    )

    col1, col2 = st.columns(2)

    with col1:
        st.subheader("✅ Strengths")

        strengths = result.get("strengths", [])

        if strengths:
            for strength in strengths:
                st.markdown(f"- {strength}")
        else:
            st.write("No strengths returned.")

    with col2:
        st.subheader("🚨 Critical Issues")

        issues = result.get("critical_issues", [])

        if issues:
            for issue in issues:
                st.markdown(f"- {issue}")
        else:
            st.success("No critical issues identified.")

    st.divider()

    st.subheader("🛠️ Recommended Improvements")
    render_improvements(
        result.get("improvements", [])
    )

    st.divider()

    # Keyword analysis
    st.subheader("🔑 Keyword Analysis")

    keyword_data = result.get(
        "keyword_analysis",
        {},
    )

    found = keyword_data.get(
        "important_keywords_found",
        [],
    )

    missing = keyword_data.get(
        "important_keywords_missing",
        [],
    )

    keyword_col1, keyword_col2 = st.columns(2)

    with keyword_col1:
        st.markdown("**Keywords found**")

        if found:
            st.write(", ".join(found))
        else:
            st.caption("No keywords were identified.")

    with keyword_col2:
        st.markdown("**Potentially missing keywords**")

        if missing:
            st.write(", ".join(missing))
        else:
            st.success("No important missing keywords identified.")

    keyword_note = keyword_data.get("note", "")

    if keyword_note:
        st.caption(keyword_note)

    st.divider()

    # Missing sections
    st.subheader("📋 Missing Sections")

    missing_sections = result.get(
        "missing_sections",
        [],
    )

    if missing_sections:
        for section in missing_sections:
            st.markdown(f"- {section}")
    else:
        st.success("No obvious missing sections identified.")

    st.divider()

    # Bullet improvements
    st.subheader("✍️ Bullet Point Improvements")

    bullet_improvements = result.get(
        "bullet_point_improvements",
        [],
    )

    if bullet_improvements:
        for item in bullet_improvements:
            original = item.get("original", "")
            improved = item.get("improved", "")
            reason = item.get("reason", "")

            with st.expander(
                original[:100] if original else "Bullet improvement"
            ):
                st.markdown("**Original**")
                st.write(original)

                st.markdown("**Improved**")
                st.success(improved)

                if reason:
                    st.caption(reason)
    else:
        st.info(
            "No bullet-point rewrites were suggested."
        )

    st.divider()

    # Formatting checklist
    st.subheader("🎨 ATS Formatting Checklist")

    formatting_items = result.get(
        "formatting_checklist",
        [],
    )

    if formatting_items:
        for item in formatting_items:
            name = item.get("item", "Formatting item")
            status = item.get("status", "Unknown")
            recommendation = item.get(
                "recommendation",
                "",
            )

            st.markdown(
                f"**{name}:** {status}"
            )

            if recommendation:
                st.caption(recommendation)

    st.divider()

    # Raw JSON download
    st.subheader("⬇️ Export Analysis")

    json_data = json.dumps(
        result,
        indent=2,
        ensure_ascii=False,
    )

    st.download_button(
        label="Download analysis as JSON",
        data=json_data,
        file_name="resume_analysis.json",
        mime="application/json",
    )
